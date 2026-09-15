import boto3
import csv
import io

from urllib.parse import unquote_plus

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from botocore.exceptions import ClientError


s3 = boto3.client("s3")


# ---------------------------------------------------------
# Configuration
# ---------------------------------------------------------

INPUT_FILE = "input/student_marks.txt"
OUTPUT_FILE = "output/duplicate_report.xlsx"


def lambda_handler(event, context):

    # -----------------------------------------------------
    # 1. Get bucket and control-file information
    # -----------------------------------------------------

    record = event["Records"][0]

    bucket_name = record["s3"]["bucket"]["name"]

    raw_key = record["s3"]["object"]["key"]

    control_file_key = unquote_plus(raw_key)

    print(
        f"Control file received: "
        f"s3://{bucket_name}/{control_file_key}"
    )

    # -----------------------------------------------------
    # 2. Make sure the event is for a .ctl file
    # -----------------------------------------------------

    if not control_file_key.lower().endswith(".ctl"):

        print("Uploaded file is not a .ctl file. Skipping.")

        return {
            "statusCode": 200,
            "message": "File skipped because it is not a .ctl file."
        }

    # -----------------------------------------------------
    # 3. Check that the input TXT file exists
    # -----------------------------------------------------

    print(
        f"Checking input file: "
        f"s3://{bucket_name}/{INPUT_FILE}"
    )

    try:

        s3.head_object(
            Bucket=bucket_name,
            Key=INPUT_FILE
        )

    except ClientError as error:

        error_code = error.response["Error"]["Code"]

        if error_code in ("404", "NoSuchKey", "NotFound"):

            print("Input TXT file was not found.")

            return {
                "statusCode": 400,
                "message": "Input TXT file is missing."
            }

        # Access denied, throttling, KMS failure, etc.
        # These are NOT "file missing" and must not be reported as such.
        print(f"Error checking {INPUT_FILE}: {error_code}")

        raise

    # -----------------------------------------------------
    # 4. Read TXT file from S3
    # -----------------------------------------------------

    response = s3.get_object(
        Bucket=bucket_name,
        Key=INPUT_FILE
    )

    file_content = response["Body"].read().decode("utf-8")

    print("Input TXT file downloaded successfully.")

    # -----------------------------------------------------
    # 5. Parse pipe-delimited TXT file
    # -----------------------------------------------------

    reader = csv.reader(
        io.StringIO(file_content),
        delimiter="|"
    )

    rows = list(reader)

    if not rows:

        print("Input TXT file is empty.")

        return {
            "statusCode": 400,
            "message": "Input TXT file is empty."
        }

    # -----------------------------------------------------
    # 6. Separate header and data
    # -----------------------------------------------------

    header = rows[0]

    data_rows = rows[1:]

    print(
        f"Total records read: "
        f"{len(data_rows)}"
    )

    # -----------------------------------------------------
    # 7. Find unique records and duplicates
    # -----------------------------------------------------

    seen = set()

    clean_records = []

    duplicate_records = []

    # Count how many times each record appears
    record_counts = {}

    for row in data_rows:

        row_tuple = tuple(row)

        record_counts[row_tuple] = (
            record_counts.get(row_tuple, 0) + 1
        )

    # Keep first occurrence.
    # Every later occurrence is treated as a duplicate.
    for row in data_rows:

        row_tuple = tuple(row)

        if row_tuple in seen:

            duplicate_records.append(row)

        else:

            seen.add(row_tuple)

            clean_records.append(row)

    # -----------------------------------------------------
    # 8. Calculate duplicate count
    # -----------------------------------------------------

    total_duplicates = len(duplicate_records)

    print(
        f"Unique records: "
        f"{len(clean_records)}"
    )

    print(
        f"Duplicate records: "
        f"{total_duplicates}"
    )

    # -----------------------------------------------------
    # 9. Create Excel workbook
    # -----------------------------------------------------

    workbook = Workbook()

    # -----------------------------------------------------
    # Sheet 1 - Clean Records
    # -----------------------------------------------------

    clean_sheet = workbook.active

    clean_sheet.title = "Clean Records"

    # Write header
    clean_sheet.append(header)

    # Write clean records
    for row in clean_records:

        clean_sheet.append(row)

    # -----------------------------------------------------
    # Add filter to clean records
    #
    # Note: the filter range must be set BEFORE the Total
    # Marks row is added, so the filter covers data rows
    # only. Filtering a totals row into the data is wrong.
    # -----------------------------------------------------

    last_clean_row = clean_sheet.max_row

    last_column_letter = (
        clean_sheet.cell(
            row=1,
            column=clean_sheet.max_column
        ).column_letter
    )

    clean_sheet.auto_filter.ref = (
        f"A1:{last_column_letter}{last_clean_row}"
    )

    # Freeze header
    clean_sheet.freeze_panes = "A2"

    # -----------------------------------------------------
    # Calculate Total Marks
    # -----------------------------------------------------

    marks_column = None

    for column_number, column_name in enumerate(
        header,
        start=1
    ):

        if column_name.strip().lower() == "marks":

            marks_column = column_number

            break

    total_marks = 0

    if marks_column is not None:

        for row in clean_records:

            try:

                total_marks += float(
                    row[marks_column - 1]
                )

            except (ValueError, TypeError):

                pass

    print(
        f"Total marks of clean records: "
        f"{total_marks}"
    )

    # Whole numbers read better as 1716 than 1716.0, both in
    # the Excel cell and in the JSON response.
    total_marks_display = (
        int(total_marks)
        if float(total_marks).is_integer()
        else total_marks
    )

    # -----------------------------------------------------
    # Add Total Marks row
    # -----------------------------------------------------

    total_row = None

    if marks_column is None:

        print(
            "No 'Marks' column found in the header. "
            "Skipping the Total Marks row."
        )

    else:

        total_row = clean_sheet.max_row + 1

        # The label goes in the column immediately LEFT of Marks and the
        # total goes IN the Marks column.
        #
        # Writing both to the same column would make the number overwrite
        # the label, and the sheet would show a bare 1716 with no label.
        label_column = (
            marks_column - 1
            if marks_column > 1
            else marks_column + 1
        )

        clean_sheet.cell(
            row=total_row,
            column=label_column,
            value="Total Marks"
        )

        clean_sheet.cell(
            row=total_row,
            column=marks_column,
            value=total_marks_display
        )

    # -----------------------------------------------------
    # Format Total Marks row
    # -----------------------------------------------------

    blue_fill = PatternFill(
        fill_type="solid",
        fgColor="5B9BD5"
    )

    bold_font = Font(
        bold=True
    )

    if total_row is not None:

        for column_number in range(
            1,
            clean_sheet.max_column + 1
        ):

            cell = clean_sheet.cell(
                row=total_row,
                column=column_number
            )

            cell.fill = blue_fill

            cell.font = bold_font

    # -----------------------------------------------------
    # Sheet 2 - Duplicate Report
    # -----------------------------------------------------

    duplicate_sheet = workbook.create_sheet(
        "Duplicate Report"
    )

    # Title
    duplicate_sheet["A1"] = "Duplicate Report"

    duplicate_sheet["A1"].font = Font(
        bold=True,
        size=14
    )

    # Total duplicate count
    duplicate_sheet["A2"] = (
        "Total Duplicate Records"
    )

    duplicate_sheet["B2"] = total_duplicates

    duplicate_sheet["A2"].font = Font(
        bold=True
    )

    duplicate_sheet["B2"].font = Font(
        bold=True
    )

    # -----------------------------------------------------
    # Duplicate table starts at row 4
    # -----------------------------------------------------

    duplicate_header_row = 4

    # Add original columns
    for column_number, column_name in enumerate(
        header,
        start=1
    ):

        duplicate_sheet.cell(
            row=duplicate_header_row,
            column=column_number,
            value=column_name
        )

    # Add Duplicate Count column
    duplicate_count_column = (
        len(header) + 1
    )

    duplicate_sheet.cell(
        row=duplicate_header_row,
        column=duplicate_count_column,
        value="Duplicate Count"
    )

    # Bold the duplicate table header
    for column_number in range(
        1,
        duplicate_count_column + 1
    ):

        duplicate_sheet.cell(
            row=duplicate_header_row,
            column=column_number
        ).font = Font(bold=True)

    # -----------------------------------------------------
    # Write duplicate records
    # -----------------------------------------------------

    duplicate_row_number = 5

    for row in duplicate_records:

        for column_number, value in enumerate(
            row,
            start=1
        ):

            duplicate_sheet.cell(
                row=duplicate_row_number,
                column=column_number,
                value=value
            )

        # Show how many times the complete record
        # appears in the original input
        row_tuple = tuple(row)

        duplicate_sheet.cell(
            row=duplicate_row_number,
            column=duplicate_count_column,
            value=record_counts[row_tuple]
        )

        duplicate_row_number += 1

    # -----------------------------------------------------
    # Add filter to duplicate table
    # -----------------------------------------------------

    last_duplicate_row = (
        duplicate_sheet.max_row
    )

    last_duplicate_column_letter = (
        duplicate_sheet.cell(
            row=duplicate_header_row,
            column=duplicate_count_column
        ).column_letter
    )

    duplicate_sheet.auto_filter.ref = (
        f"A{duplicate_header_row}:"
        f"{last_duplicate_column_letter}"
        f"{last_duplicate_row}"
    )

    # Freeze duplicate table header
    duplicate_sheet.freeze_panes = "A5"

    # -----------------------------------------------------
    # Basic column widths
    # -----------------------------------------------------

    for worksheet in [
        clean_sheet,
        duplicate_sheet
    ]:

        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = (
                column_cells[0].column_letter
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                30
            )

    # -----------------------------------------------------
    # 10. Save Excel workbook in memory
    # -----------------------------------------------------

    excel_buffer = io.BytesIO()

    workbook.save(excel_buffer)

    excel_buffer.seek(0)

    # -----------------------------------------------------
    # 11. Upload Excel file to S3
    # -----------------------------------------------------

    s3.put_object(
        Bucket=bucket_name,
        Key=OUTPUT_FILE,
        Body=excel_buffer.getvalue(),
        ContentType=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    print(
        f"Excel file created successfully: "
        f"s3://{bucket_name}/{OUTPUT_FILE}"
    )

    # -----------------------------------------------------
    # 12. Return Lambda response
    # -----------------------------------------------------

    return {
        "statusCode": 200,
        "input_file": INPUT_FILE,
        "total_input_records": len(data_rows),
        "clean_records": len(clean_records),
        "duplicate_records": total_duplicates,
        "total_marks": total_marks_display,
        "output_file": OUTPUT_FILE,
        "message": (
            "Duplicate removal and Excel reporting "
            "completed successfully."
        )
    }
